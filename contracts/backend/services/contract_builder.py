from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from backend.services.calculator import calculate_services, format_money, load_tariffs
from backend.services.dates import contract_date_ru
from backend.services.formatting import clean_text, html_escape
from backend.services.number_to_words import rubles_to_words


DEMO_NOTICE = (
    "Synthetic portfolio template. Replace every legal clause and requisite "
    "after review by qualified counsel before any real use."
)


def fixed_sections_for(data: dict[str, Any]) -> list[tuple[str, list[str]]]:
    return [
        ("1. ДЕМОНСТРАЦИОННЫЙ ПРЕДМЕТ", [
            "1.1. Исполнитель оказывает демонстрационные услуги по ведению карточек организации в картографических сервисах.",
            f"1.2. Демонстрационный срок составляет {data['months']} месяцев.",
        ]),
        ("2. ДЕМОНСТРАЦИОННАЯ СТОИМОСТЬ", [
            f"2.1. Синтетическая стоимость составляет {format_money(data['total_price'])}.",
            "2.2. Этот текст не является офертой или готовым юридическим документом.",
        ]),
        ("3. КОНФИДЕНЦИАЛЬНОСТЬ", [
            "3.1. Реальные персональные данные и реквизиты должны обрабатываться только в разрешённой защищённой среде.",
            "3.2. Сгенерированные документы и исходные файлы не должны попадать в публичный репозиторий.",
        ]),
    ]


def normalize_payload(payload: dict[str, Any]) -> dict[str, Any]:
    data = dict(payload or {})
    data["contract_number"] = clean_text(data.get("contract_number")) or "DEMO-001"
    data["contract_date_formatted"] = contract_date_ru(data.get("contract_date"))
    data["months"] = int(data.get("months") or 3)
    data["total_price"] = int(float(str(data.get("total_price") or 0).replace(" ", "")))
    data["include_reputation"] = bool(data.get("include_reputation", True))
    reputation = data.get("reputation_total")
    data["reputation_total"] = None if reputation in (None, "") else int(reputation)
    data.update(calculate_services(
        data["months"], data["total_price"], data["reputation_total"], data["include_reputation"]
    ))
    data["total_price_text"] = rubles_to_words(data["total_price"])
    return data


def validate_payload(data: dict[str, Any]) -> list[str]:
    required = {
        "customer_legal_name": "название заказчика",
        "customer_long_name": "представитель заказчика",
        "customer_inn": "ИНН заказчика",
        "org_name": "название демонстрационного объекта",
        "org_address": "адрес демонстрационного объекта",
    }
    return [f"Не заполнено поле: {label}" for key, label in required.items() if not clean_text(data.get(key))]


def vat_phrase(data: dict[str, Any]) -> str:
    mode = clean_text(data.get("vat_mode"))
    return ", включая НДС 5%." if mode == "vat_5" else ", без НДС."


def payment_lines(data: dict[str, Any]) -> list[str]:
    schedule = data.get("payment_schedule") or []
    if schedule:
        return [clean_text(item) for item in schedule if clean_text(item)]
    return ["Демонстрационный платёжный график формируется после юридической проверки."]


def requisites_block(data: dict[str, Any]) -> str:
    fields = [
        ("Заказчик", "customer_legal_name"),
        ("ИНН", "customer_inn"),
        ("КПП", "customer_kpp"),
        ("ОГРН", "customer_ogrn"),
        ("Адрес", "customer_legal_address"),
        ("Расчетный счет", "customer_rs"),
        ("Банк", "customer_bank"),
        ("БИК", "customer_bik"),
        ("Корреспондентский счет", "customer_ks"),
    ]
    return "\n".join(f"{label}: {clean_text(data.get(key))}" for label, key in fields)


def executor_block() -> str:
    return "\n".join(str(item) for item in load_tariffs()["executor"]["details"])


def contract_filename(data: dict[str, Any], extension: str) -> str:
    number = clean_text(data.get("contract_number")) or "DEMO-001"
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in number)
    return f"synthetic_contract_{safe}.{extension}"


def _service_rows(data: dict[str, Any]) -> str:
    rows = []
    for item in data["rows"]:
        rows.append(
            f"<tr><td>{html_escape(item['name'])}</td><td>{format_money(item['amount'])}</td></tr>"
        )
    return "".join(rows)


def build_preview_html(payload: dict[str, Any]) -> str:
    data = normalize_payload(payload)
    warnings = validate_payload(data)
    warning_html = "".join(f"<li>{html_escape(item)}</li>" for item in warnings)
    sections = "".join(
        f"<h2>{html_escape(title)}</h2>" + "".join(f"<p>{html_escape(text)}</p>" for text in paragraphs)
        for title, paragraphs in fixed_sections_for(data)
    )
    return f"""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<title>Синтетический договор {html_escape(data['contract_number'])}</title>
<style>body{{font-family:Arial;max-width:900px;margin:2rem auto;line-height:1.45}}table{{border-collapse:collapse;width:100%}}td,th{{border:1px solid #999;padding:.5rem}}.notice{{background:#fff3cd;padding:1rem}}</style>
</head><body><p class="notice"><b>{html_escape(DEMO_NOTICE)}</b></p>
<h1>ДЕМОНСТРАЦИОННЫЙ ДОГОВОР №{html_escape(data['contract_number'])}</h1>
<p>{html_escape(data.get('customer_legal_name'))} и {html_escape(load_tariffs()['executor']['name'])}.</p>
{f'<ul>{warning_html}</ul>' if warning_html else ''}{sections}
<h2>Расчёт</h2><table><tr><th>Позиция</th><th>Сумма</th></tr>{_service_rows(data)}</table>
<h2>Синтетические реквизиты</h2><pre>{html_escape(requisites_block(data))}</pre>
<pre>{html_escape(executor_block())}</pre></body></html>"""


def build_docx(payload: dict[str, Any]) -> bytes:
    data = normalize_payload(payload)
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(f"ДЕМОНСТРАЦИОННЫЙ ДОГОВОР №{data['contract_number']}")
    run.bold = True
    run.font.size = Pt(14)

    notice = document.add_paragraph(DEMO_NOTICE)
    notice.runs[0].bold = True
    document.add_paragraph(
        f"{clean_text(data.get('customer_legal_name'))} и {load_tariffs()['executor']['name']}."
    )

    for title, paragraphs in fixed_sections_for(data):
        document.add_heading(title, level=1)
        for text in paragraphs:
            document.add_paragraph(text)

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Позиция"
    table.rows[0].cells[1].text = "Сумма"
    for item in data["rows"]:
        cells = table.add_row().cells
        cells[0].text = item["name"]
        cells[1].text = format_money(item["amount"])

    document.add_heading("Синтетические реквизиты", level=1)
    document.add_paragraph(requisites_block(data))
    document.add_paragraph(executor_block())

    output = BytesIO()
    document.save(output)
    return output.getvalue()

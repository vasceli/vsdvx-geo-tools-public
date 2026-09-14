from __future__ import annotations

import re
import tempfile
from pathlib import Path
from typing import Any

import fitz  # PyMuPDF
from docx import Document
from openpyxl import load_workbook

from backend.services.formatting import clean_text, initials_from_name

CUSTOMER_FIELDS = [
    "customer_legal_name", "customer_long_name", "customer_director_position", "customer_basis",
    "customer_short_sign", "customer_inn", "customer_kpp", "customer_ogrn", "customer_rs",
    "customer_ks", "customer_bik", "customer_bank", "customer_bank_address", "customer_legal_address",
]


def _read_pdf(path: Path) -> str:
    parts = []
    with fitz.open(path) as doc:
        for page in doc:
            parts.append(page.get_text("text"))
    return "\n".join(parts)


def _read_docx(path: Path) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    wb = load_workbook(path, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            line = " | ".join(clean_text(c) for c in row if c is not None)
            if line:
                parts.append(line)
    return "\n".join(parts)


def extract_text_from_file(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(content)
        path = Path(tmp.name)
    try:
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if suffix in (".xlsx", ".xlsm"):
            return _read_xlsx(path)
        return content.decode("utf-8", errors="ignore")
    finally:
        try:
            path.unlink(missing_ok=True)
        except Exception:
            pass


def _first(patterns: list[str], text: str, flags: int = re.I | re.S) -> str:
    for pattern in patterns:
        m = re.search(pattern, text, flags)
        if m:
            return clean_text(m.group(1))
    return ""


def _lines(text: str) -> list[str]:
    return [clean_text(line) for line in text.replace("\u00a0", " ").replace("\u200b", "").splitlines() if clean_text(line)]


def _line_index(lines: list[str], label: str) -> int:
    label_re = re.compile(label, re.I)
    for i, line in enumerate(lines):
        if label_re.search(line):
            return i
    return -1


def _value_after_label_lines(lines: list[str], label_pattern: str, stop_label_patterns: list[str] | None = None) -> str:
    """Extracts values from bank requisites where PDFs put label and value on different lines."""
    idx = _line_index(lines, label_pattern)
    if idx < 0:
        return ""
    label_re = re.compile(label_pattern, re.I)
    same_line = clean_text(label_re.sub("", lines[idx], count=1).strip(" :-№"))
    if same_line and not re.fullmatch(r"[A-ЯЁа-яё .\-/]+", same_line):
        return same_line

    stop_defaults = [
        r"^ИНН(?:\s+банка)?$", r"^КПП$", r"^ОГРН", r"^Расчетный\s+счет$", r"^Р/счет$",
        r"^Банк$", r"^Юридический\s+адрес\s+банка$", r"^Корр\.\s*счет\s+банка$",
        r"^БИК\s+банка$", r"^Директор$", r"^Юридический\s+адрес$", r"^Фактический\s+адрес$",
    ]
    stop_patterns = [re.compile(p, re.I) for p in (stop_label_patterns or stop_defaults)]
    values: list[str] = []
    for line in lines[idx + 1:]:
        if any(p.search(line) for p in stop_patterns):
            break
        # Stop at the next obvious field name, not at meaningful address/company text.
        if re.search(r"^(Полное\s+наименование|Реквизиты\s+счета)", line, re.I):
            break
        values.append(line)
        # Most scalar bank values are a single line.
        if re.search(r"(\d{20}|\d{9}|\d{10,12}|\d{13,15}|[АA]О\s+[\"«])", line):
            break
    return clean_text(" ".join(values))


def _label_value(patterns: list[str], text: str, lines: list[str]) -> str:
    value = _first(patterns, text)
    if value:
        return value
    for p in patterns:
        # Convert simple regex label before first capture into line lookup when possible.
        m = re.match(r"\(?\?:?(.+?)\)??\\s\*\[:", p)
        if m:
            value = _value_after_label_lines(lines, m.group(1))
            if value:
                return value
    return ""


def _customer_scope(text: str) -> str:
    """Prefer customer requisites block when a whole contract is uploaded."""
    candidates = []
    for m in re.finditer(
        r"Заказчик\s+\n?\s*Исполнитель\s+(.*?)(?:\n\s*_________________|\Z)",
        text,
        flags=re.I | re.S,
    ):
        block = m.group(1)
        if "ИНН" in block and ("ОГРН" in block or "ОГРНИП" in block):
            candidates.append(block)
    return candidates[-1] if candidates else text


def _bank_after_account(scope: str) -> str:
    m = re.search(
        r"(?:р\s*/\s*с|Р/счет|расч[её]тный\s+сч[её]т)\s*[:№]?\s*\d{20}\s*(?:\n|\r\n?)(.*?)(?:\n|\r\n?)\s*БИК",
        scope,
        flags=re.I | re.S,
    )
    if not m:
        return ""
    bank = clean_text(m.group(1))
    bank = re.sub(r"\s+", " ", bank)
    return bank.strip(" ,")


def _legal_address(scope: str, lines: list[str]) -> str:
    addr = _first([
        r"Юридическ(?:ий|ий)\s+адрес\s*[:\-]?\s*(.*?)(?=\n\s*(?:Фактический\s+адрес|р\s*/\s*с|Р/счет|расч|БИК|Кор|К/с|ИНН|ОГРН|КПП|_________________|$))",
        r"Адрес\s*[:\-]?\s*([^\n\r]+)",
    ], scope)
    if addr and not re.search(r"^(Фактический|ОБЩЕСТВО|ООО|ИП)\b", addr, re.I):
        return re.sub(r"\s+", " ", addr).strip()

    # T-Bank PDFs list labels first, then values. Extract first address block after director.
    start = -1
    for i, line in enumerate(lines):
        if re.match(r"^[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+$", line):
            start = i + 1
            break
    if start >= 0:
        addr_lines = []
        for line in lines[start:]:
            if re.search(r"^(ИНН|КПП|ОГРН|Расчетный счет|Банк|Юридический адрес банка|Корр\.|БИК)", line, re.I):
                break
            # Stop at factual address: second postal-code block.
            if addr_lines and re.match(r"^\d{6}\b", line):
                break
            if re.match(r"^\d{6}\b", line) or addr_lines:
                addr_lines.append(line)
        if addr_lines:
            return re.sub(r"\s+", " ", " ".join(addr_lines)).strip()
    return ""


def _parse_tbank_style(text: str, result: dict[str, Any]) -> dict[str, Any]:
    """Handles TBank requisites PDFs whose text layer is visually ordered, not semantically ordered."""
    lines = _lines(text)
    if not any("Tbank" in line or "ТБанк" in line for line in lines):
        return result

    # Legal name is commonly split into two lines.
    legal_name = _first([
        r"(ОБЩЕСТВО\s+С\s+ОГРАНИЧЕННОЙ\s+ОТВЕТСТВЕННОСТЬЮ\s+[\"«][^\"»]+[\"»])",
        r"(ОБЩЕСТВО\s+С\s+ОГРАНИЧЕННОЙ\s+\n?\s*ОТВЕТСТВЕННОСТЬЮ\s+[\"«][^\"»]+[\"»])",
    ], text)
    if legal_name:
        result["customer_legal_name"] = re.sub(r"\s+", " ", legal_name).strip()

    director = _first([r"\n\s*([А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+)\s*\n\s*\d{6}"], text)
    if not director:
        for line in lines:
            if re.match(r"^[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+$", line):
                director = line
                break
    if director:
        result["customer_long_name"] = director
        result["customer_director_position"] = "директора"
        result["customer_short_sign"] = initials_from_name(director)

    result["customer_legal_address"] = _legal_address(text, lines) or result.get("customer_legal_address", "")

    # Scalar fields. Prefer line-based extraction because labels and values are often split by newlines.
    mappings = {
        "customer_inn": r"^ИНН$",
        "customer_kpp": r"^КПП$",
        "customer_ogrn": r"^ОГРН",
        "customer_rs": r"^Расчетный\s+счет$",
        "customer_bank": r"^Банк$",
        "customer_bank_address": r"^Юридический\s+адрес\s+банка$",
        "customer_ks": r"^Корр\.\s*счет\s+банка$",
        "customer_bik": r"^БИК\s+банка$",
    }
    for key, label in mappings.items():
        val = _value_after_label_lines(lines, label)
        if val:
            result[key] = val

    # Regex fallback for numeric lines in case the line-based helper got conservative.
    flat = "\n".join(lines)
    result["customer_inn"] = result.get("customer_inn") or _first([r"ИНН\s*(\d{10,12})"], flat)
    result["customer_kpp"] = result.get("customer_kpp") or _first([r"КПП\s*(\d{9})"], flat)
    result["customer_ogrn"] = result.get("customer_ogrn") or _first([r"ОГРН\s*(\d{13,15})"], flat)
    result["customer_rs"] = result.get("customer_rs") or _first([r"Расчетный\s+счет\s*(\d{20})"], flat)
    result["customer_ks"] = result.get("customer_ks") or _first([r"Корр\.\s*счет\s+банка\s*(\d{20})"], flat)
    result["customer_bik"] = result.get("customer_bik") or _first([r"БИК\s+банка\s*(\d{8,9})"], flat)
    return result


def parse_requisites_text(text: str) -> dict[str, Any]:
    raw = text
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+", " ", text)
    flat_text = re.sub(r"\s+", " ", text)
    scope = _customer_scope(text)
    flat_scope = re.sub(r"\s+", " ", scope)
    scope_lines = _lines(scope)

    legal_name = _first([
        r"((?:ОБЩЕСТВО|Общество)\s+с\s+ограниченной\s+ответственностью\s+[«\"“][^»\"”]+[»\"”])",
        r"(?:^|\n)\s*((?:ООО|ИП)\s+[«\"“]?[^\n\r|]+[»\"”]?)",
        r"(?:Наименование|Организация|Компания|Полное\s+наименование\s+организации)\s*[:\-]?\s*([^\n\r|]+)",
    ], scope) or _first([
        r"((?:ОБЩЕСТВО|Общество)\s+с\s+ограниченной\s+ответственностью\s+[«\"“][^»\"”]+[»\"”])",
        r"(?:^|\n)\s*((?:ООО|ИП)\s+[«\"“]?[^\n\r|]+[»\"”]?)",
    ], text)

    customer_long_name = _first([
        r"в лице\s+(?:генерального директора|директора)\s+([А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+)",
        r"(?:Генеральный директор|Директор|Руководитель)\s*[:\-]?\s*([А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+)",
    ], flat_text)

    director_position = "генерального директора" if re.search(r"генеральн", flat_text, re.I) else "директора"
    bank_after_account = _bank_after_account(scope)

    result = {
        "customer_legal_name": legal_name,
        "customer_long_name": customer_long_name,
        "customer_director_position": director_position,
        "customer_basis": "Устава",
        "customer_short_sign": initials_from_name(customer_long_name),
        "customer_inn": _first([r"ИНН\s*[:№]?\s*(\d{10,12})"], scope),
        "customer_kpp": _first([r"КПП\s*[:№]?\s*(\d{9})"], scope),
        "customer_ogrn": _first([r"ОГРН(?:ИП)?\s*[:№]?\s*(\d{13,15})"], scope),
        "customer_rs": _first([r"(?:р\s*/\s*с|р/счет|расч[её]тный\s+сч[её]т|Р/счет|Расчетный\s+счет)\s*[:№]?\s*(\d{20})"], scope),
        "customer_ks": _first([r"(?:к\s*/\s*с|кор\.?\s*сч[её]т(?:\s+банка)?|корреспондентский\s+сч[её]т|К/с)\s*[:№]?\s*(\d{20})"], scope),
        "customer_bik": _first([r"БИК(?:\s+банка)?\s*[:№]?\s*(\d{8,9})"], scope),
        "customer_bank": bank_after_account or _value_after_label_lines(scope_lines, r"^Банк$") or _first([
            r"(?:^|\n)\s*(?:Банк|Наименование банка)\s*[:\-]\s*([^\n\r|]+)",
            r"(?:в\s+)(АО\s+[«\"“][^»\"”]+[»\"”])",
            r"((?:ПАО|АО)\s+[«\"“][^»\"”]+[»\"”])",
        ], scope),
        "customer_bank_address": _value_after_label_lines(scope_lines, r"^Юридический\s+адрес\s+банка$") or _value_after_label_lines(scope_lines, r"^Адрес\s+банка$") or _first([
            r"Юридический\s+адрес\s+банка\s*[:\-]?\s*([^\n\r]+)",
            r"Адрес\s+банка\s*[:\-]?\s*([^\n\r]+)",
        ], scope),
        "customer_legal_address": _legal_address(scope, scope_lines),
    }

    result = _parse_tbank_style(text, result)

    if not result.get("customer_short_sign"):
        sign = _first([r"_________________\s*([^\n\r]+)"], scope)
        if sign:
            result["customer_short_sign"] = sign

    for key, value in list(result.items()):
        if isinstance(value, str):
            value = re.sub(r"\s+", " ", value).strip()
            if "{{" in value or "}}" in value:
                value = ""
            result[key] = value
    if len(result.get("customer_bik", "")) == 8:
        result["customer_bik"] = "0" + result["customer_bik"]

    result["raw_text_preview"] = clean_text(raw[:3000])
    result["missing_fields"] = [
        k for k, v in result.items()
        if k.startswith("customer_") and not v and k not in {"customer_kpp", "customer_bank_address"}
    ]
    return result


def parse_requisites_file(filename: str, content: bytes) -> dict[str, Any]:
    text = extract_text_from_file(filename, content)
    return parse_requisites_text(text)
